"""Leitura das fontes: exportação do WhatsApp Business, exportação do Instagram e planilhas de prospecção.

Nenhuma fonte é acessada pela internet: a Renata exporta as conversas pelo próprio app e coloca os arquivos
na pasta de dados. Assim o agente não toca na conta (que já foi restrita por envio em sequência) e nada sai
do computador além do texto enviado para análise.
"""

from __future__ import annotations

import csv
import io
import json
import re
import unicodedata
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path

from .modelos import Lead, Mensagem

# 16/10/2026 14:32 - Nome: texto   (Android)
_WA_ANDROID = re.compile(r"^(\d{1,2}/\d{1,2}/\d{2,4}),? (\d{1,2}:\d{2})(?::\d{2})? - ([^:]+?): (.*)$")
# [16/10/2026, 14:32:10] Nome: texto   (iPhone)
_WA_IOS = re.compile(r"^\[(\d{1,2}/\d{1,2}/\d{2,4}),? (\d{1,2}:\d{2})(?::\d{2})?\] ([^:]+?): (.*)$")
# Linhas de sistema do WhatsApp (sem autor) começam com data mas não casam com os padrões acima
_WA_INICIO_DATA = re.compile(r"^\[?\d{1,2}/\d{1,2}/\d{2,4}")
_INVISIVEIS = dict.fromkeys(map(ord, "‎‏‪‬﻿"), None)
_MIDIA = re.compile(r"<(m[íi]dia oculta|arquivo de m[íi]dia oculto|media omitted)>", re.I)


def normalizar(texto: str) -> str:
    sem_acento = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", sem_acento).strip().lower()


def so_digitos(texto: str) -> str:
    return re.sub(r"\D", "", texto or "")


def chave_telefone(telefone: str) -> str:
    """Últimos 10 dígitos (DDD + número sem o 9 extra fica estável entre formatos)."""
    d = so_digitos(telefone)
    if len(d) < 10:
        return ""
    if d.startswith("55") and len(d) > 11:
        d = d[2:]
    # 31 9 8888-7777 e 31 8888-7777 viram a mesma chave
    if len(d) == 11 and d[2] == "9":
        d = d[:2] + d[3:]
    return "tel:" + d[-10:]


def _data(dia: str, hora: str) -> datetime:
    d, m, a = dia.split("/")
    ano = int(a) + 2000 if len(a) == 2 else int(a)
    h, mi = hora.split(":")
    return datetime(ano, int(m), int(d), int(h), int(mi))


def _nome_da_conversa(arquivo: str) -> str:
    base = Path(arquivo).stem
    for prefixo in ("Conversa do WhatsApp com ", "WhatsApp Chat with ", "WhatsApp Chat - "):
        if base.startswith(prefixo):
            return base[len(prefixo):].strip()
    return base.strip()


def _ler_textos_whatsapp(pasta: Path) -> list[tuple[str, str]]:
    """Devolve (nome_do_arquivo, conteúdo) de cada .txt solto ou dentro de .zip (só .txt é lido do zip)."""
    textos = []
    for caminho in sorted(pasta.rglob("*")):
        if caminho.suffix.lower() == ".txt":
            textos.append((caminho.name, caminho.read_text(encoding="utf-8", errors="replace")))
        elif caminho.suffix.lower() == ".zip":
            with zipfile.ZipFile(caminho) as z:
                for info in z.infolist():
                    if info.filename.lower().endswith(".txt") and not info.is_dir():
                        conteudo = z.read(info).decode("utf-8", errors="replace")
                        nome = caminho.name if Path(info.filename).name == "_chat.txt" else Path(info.filename).name
                        textos.append((nome, conteudo))
    return textos


def _parse_whatsapp(conteudo: str) -> list[tuple[datetime, str, str]]:
    mensagens: list[tuple[datetime, str, str]] = []
    for linha in conteudo.splitlines():
        linha = linha.translate(_INVISIVEIS).rstrip()
        achou = _WA_ANDROID.match(linha) or _WA_IOS.match(linha)
        if achou:
            dia, hora, autor, texto = achou.groups()
            texto = _MIDIA.sub("[mídia]", texto)
            mensagens.append((_data(dia, hora), autor.strip(), texto))
        elif mensagens and linha and not _WA_INICIO_DATA.match(linha):
            quando, autor, texto = mensagens[-1]
            mensagens[-1] = (quando, autor, texto + "\n" + linha)
    return mensagens


def _identificar_nossos(autores_por_conversa: list[set[str]], nossos_nomes: list[str]) -> set[str]:
    nossos = {normalizar(n) for n in nossos_nomes}
    if len(autores_por_conversa) >= 2:
        # quem exportou aparece em quase todas as conversas
        contagem = Counter(normalizar(a) for autores in autores_por_conversa for a in autores)
        limite = max(2, int(len(autores_por_conversa) * 0.6))
        nossos |= {a for a, n in contagem.items() if n >= limite}
    return nossos


def ler_whatsapp(pasta: Path, nossos_nomes: list[str], etiqueta: str) -> list[Lead]:
    brutos = [(arq, _parse_whatsapp(txt)) for arq, txt in _ler_textos_whatsapp(pasta)]
    brutos = [(arq, msgs) for arq, msgs in brutos if msgs]
    nossos = _identificar_nossos([{a for _, a, _ in msgs} for _, msgs in brutos], nossos_nomes)
    leads = []
    for arquivo, msgs in brutos:
        contato = _nome_da_conversa(arquivo)
        telefone = contato if len(so_digitos(contato)) >= 10 else ""
        lead = Lead(
            chave=chave_telefone(telefone) or "nome:" + normalizar(contato),
            nome=contato,
            canal="whatsapp",
            telefone=so_digitos(telefone),
            origem=arquivo,
            etiqueta=etiqueta,
        )
        for quando, autor, texto in msgs:
            lead.mensagens.append(Mensagem(quando, autor, texto, normalizar(autor) in nossos))
        lead.ordenar()
        leads.append(lead)
    return leads


def _consertar_instagram(texto: str) -> str:
    """A exportação do Instagram grava UTF-8 como se fosse Latin-1 ("NÃ£o" em vez de "Não")."""
    try:
        return texto.encode("latin-1").decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return texto


def ler_instagram(pasta: Path, nossos_nomes: list[str]) -> list[Lead]:
    nossos = {normalizar(n) for n in nossos_nomes}
    conversas: dict[str, Lead] = {}
    for arquivo in sorted(pasta.rglob("message_*.json")):
        dados = json.loads(arquivo.read_text(encoding="utf-8"))
        participantes = [_consertar_instagram(p.get("name", "")) for p in dados.get("participants", [])]
        outros = [p for p in participantes if normalizar(p) not in nossos] or participantes[-1:]
        titulo = _consertar_instagram(dados.get("title", "")) or (outros[0] if outros else arquivo.parent.name)
        caminho = dados.get("thread_path", "") or arquivo.parent.name
        perfil = Path(caminho).name.rsplit("_", 1)[0]
        chave = "ig:" + normalizar(perfil or titulo)
        lead = conversas.setdefault(
            chave,
            Lead(chave=chave, nome=titulo, canal="instagram", instagram="@" + perfil if perfil else "", origem=caminho),
        )
        for m in dados.get("messages", []):
            texto = _consertar_instagram(m.get("content", "")) or ("[mídia]" if m.get("photos") or m.get("share") else "")
            if not texto:
                continue
            autor = _consertar_instagram(m.get("sender_name", ""))
            quando = datetime.fromtimestamp(m.get("timestamp_ms", 0) / 1000)
            lead.mensagens.append(Mensagem(quando, autor, texto, normalizar(autor) in nossos))
    for lead in conversas.values():
        lead.ordenar()
    return list(conversas.values())


_COLUNAS = {
    "nome": ("clinica", "nome", "consultorio", "empresa", "contato"),
    "telefone": ("telefone", "whatsapp", "celular", "fone", "numero"),
    "instagram": ("instagram", "perfil", "@"),
    "status": ("status", "resultado", "situacao", "retorno", "resposta"),
    "lote": ("lote",),
    "data": ("data", "enviado em", "dia"),
    "observacoes": ("observ", "obs", "nota", "comentario", "especialidade", "bairro"),
    "canal": ("canal",),
}


def _mapear_colunas(cabecalho: list[str]) -> dict[str, int]:
    mapa: dict[str, int] = {}
    normalizados = [normalizar(str(c or "")) for c in cabecalho]
    for campo, chaves in _COLUNAS.items():
        for i, col in enumerate(normalizados):
            if i in mapa.values() and campo != "observacoes":
                continue
            if any(k in col for k in chaves):
                mapa.setdefault(campo, i)
    return mapa


def _linhas_planilha(caminho: Path) -> list[list]:
    if caminho.suffix.lower() == ".csv":
        texto = caminho.read_text(encoding="utf-8-sig", errors="replace")
        dialeto = csv.Sniffer().sniff(texto[:2048], delimiters=",;\t")
        return [linha for linha in csv.reader(io.StringIO(texto), dialeto)]
    from openpyxl import load_workbook  # só necessário para .xlsx

    linhas: list[list] = []
    livro = load_workbook(caminho, read_only=True, data_only=True)
    for aba in livro.worksheets:
        linhas.extend([list(l) for l in aba.iter_rows(values_only=True)])
    return linhas


def _data_planilha(valor) -> datetime | None:
    if isinstance(valor, datetime):
        return valor
    texto = str(valor or "").strip()
    for formato in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d/%m"):
        try:
            data = datetime.strptime(texto[:10], formato)
            return data.replace(year=datetime.now().year) if formato == "%d/%m" else data
        except ValueError:
            continue
    return None


_ENVIADO = ("enviad", "entregue", "sem resposta", "aguardando", "resposta automatica", "visualiz", "respond")
_NAO_ENVIADO = ("nao enviad", "pendente", "adiad", "nao aceita", "a enviar", "desconect")


def ler_planilhas(pasta: Path) -> list[Lead]:
    leads = []
    for caminho in sorted([*pasta.rglob("*.csv"), *pasta.rglob("*.xlsx")]):
        linhas = [l for l in _linhas_planilha(caminho) if any(c not in (None, "") for c in l)]
        if not linhas:
            continue
        mapa = _mapear_colunas(linhas[0])
        if "nome" not in mapa:
            continue
        lote_arquivo = re.search(r"lote[-_ ]?(\d+)", caminho.stem, re.I)
        for linha in linhas[1:]:
            def campo(nome: str) -> str:
                i = mapa.get(nome)
                return str(linha[i]).strip() if i is not None and i < len(linha) and linha[i] is not None else ""

            nome = campo("nome")
            if not nome:
                continue
            telefone = so_digitos(campo("telefone"))
            perfil = campo("instagram").lstrip("@").split("/")[-1] if campo("instagram") else ""
            canal = "instagram" if ("instagram" in normalizar(campo("canal") + caminho.stem) or (perfil and not telefone)) else "whatsapp"
            chave = chave_telefone(telefone) or ("ig:" + normalizar(perfil) if perfil else "nome:" + normalizar(nome))
            status = campo("status")
            lead = Lead(
                chave=chave,
                nome=nome,
                canal=canal,
                telefone=telefone,
                instagram="@" + perfil if perfil else "",
                origem=caminho.name,
                lote=campo("lote") or (f"Lote {lote_arquivo.group(1)}" if lote_arquivo else ""),
                observacoes=campo("observacoes"),
            )
            lead.fatos["status_planilha"] = status
            st = normalizar(status)
            data = _data_planilha(campo("data") if "data" in mapa else "")
            if data and any(k in st for k in _ENVIADO) and not any(k in st for k in _NAO_ENVIADO):
                # A planilha só registra que a mensagem 1 saiu; o texto exato não está nela.
                lead.mensagens.append(Mensagem(data.replace(hour=data.hour or 10), "Sinmed", "(mensagem de prospecção registrada na planilha; texto não salvo)", True))
            leads.append(lead)
    return leads


def mesclar(*grupos: list[Lead]) -> list[Lead]:
    """Une o mesmo contato vindo de fontes diferentes (planilha + conversa exportada)."""
    por_chave: dict[str, Lead] = {}
    por_nome: dict[str, str] = {}
    for grupo in grupos:
        for lead in grupo:
            chave = lead.chave
            if chave not in por_chave and chave.startswith("nome:") and chave[5:] in por_nome:
                chave = por_nome[chave[5:]]
            existente = por_chave.get(chave)
            if existente is None:
                por_chave[chave] = lead
                por_nome.setdefault(normalizar(lead.nome), chave)
                continue
            # a conversa real vale mais que o registro sintético da planilha
            reais = [m for m in lead.mensagens if not m.texto.startswith("(mensagem de prospecção registrada")]
            if reais:
                existente.mensagens = [
                    m for m in existente.mensagens if not m.texto.startswith("(mensagem de prospecção registrada")
                ] + reais
            existente.ordenar()
            for atributo in ("telefone", "instagram", "lote", "observacoes", "etiqueta"):
                if not getattr(existente, atributo) and getattr(lead, atributo):
                    setattr(existente, atributo, getattr(lead, atributo))
            if existente.nome.startswith(("+", "55")) or len(so_digitos(existente.nome)) >= 10:
                existente.nome = lead.nome
            existente.origem = f"{existente.origem}; {lead.origem}"
            existente.fatos.update({k: v for k, v in lead.fatos.items() if v})
    return list(por_chave.values())
